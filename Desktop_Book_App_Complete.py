import os

# --- FORCE CPU ONLY BEFORE ANY OTHER IMPORTS ---
os.environ["CUDA_VISIBLE_DEVICES"] = "-1"
os.environ["ORT_CUDA_UNAVAILABLE"] = "1"

import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog
import threading
import sys
import time
import queue
import re
import requests
import concurrent.futures

# --- DATA SCIENCE / AUDIO IMPORTS ---
import numpy as np
import sounddevice as sd
import soundfile as sf
import fitz  # PyMuPDF
from docx import Document
from pydub import AudioSegment

# --- AI IMPORTS ---
import onnxruntime as ort

try:
    from kokoro_onnx import Kokoro
except ImportError:
    Kokoro = None


# ========================================================
#             MODERN UI HELPER CLASSES
# ========================================================

class ModernMessagebox(ctk.CTkToplevel):
    def __init__(self, parent, title, message, icon="info"):
        super().__init__(parent)
        self.title(title)
        self.geometry("400x200")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()

        try:
            x = parent.winfo_x() + (parent.winfo_width() // 2) - 200
            y = parent.winfo_y() + (parent.winfo_height() // 2) - 100
            self.geometry(f"+{x}+{y}")
        except:
            pass

        text_color = "#ffffff"
        if icon == "error":
            text_color = "#ef4444"
        elif icon == "success":
            text_color = "#10b981"

        self.lbl_title = ctk.CTkLabel(self, text=title, font=("Arial", 18, "bold"), text_color=text_color)
        self.lbl_title.pack(pady=(20, 10))

        self.lbl_msg = ctk.CTkLabel(self, text=message, font=("Arial", 14), wraplength=350)
        self.lbl_msg.pack(pady=10, padx=20)

        self.btn_ok = ctk.CTkButton(self, text="OK", width=100, command=self.destroy)
        self.btn_ok.pack(pady=20)
        self.wait_window()


class ModernInputBox(ctk.CTkToplevel):
    def __init__(self, parent, title, prompt, default_val=""):
        super().__init__(parent)
        self.title(title)
        self.geometry("400x220")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()
        self.value = None

        try:
            x = parent.winfo_x() + (parent.winfo_width() // 2) - 200
            y = parent.winfo_y() + (parent.winfo_height() // 2) - 110
            self.geometry(f"+{x}+{y}")
        except:
            pass

        ctk.CTkLabel(self, text=title, font=("Arial", 18, "bold")).pack(pady=(20, 10))
        ctk.CTkLabel(self, text=prompt, font=("Arial", 14), text_color="gray").pack(pady=(0, 10))

        self.entry = ctk.CTkEntry(self, width=200, font=("Arial", 14))
        self.entry.pack(pady=10)
        if default_val: self.entry.insert(0, str(default_val))
        self.entry.focus_set()

        btn_frame = ctk.CTkFrame(self, fg_color="transparent")
        btn_frame.pack(pady=20)

        ctk.CTkButton(btn_frame, text="Cancel", fg_color="#ef4444", hover_color="#dc2626",
                      width=80, command=self.on_cancel).pack(side="left", padx=10)
        ctk.CTkButton(btn_frame, text="Confirm", fg_color="#3b82f6", hover_color="#2563eb",
                      width=80, command=self.on_confirm).pack(side="left", padx=10)

        self.bind('<Return>', lambda e: self.on_confirm())
        self.bind('<Escape>', lambda e: self.on_cancel())
        self.wait_window()

    def on_confirm(self):
        self.value = self.entry.get()
        self.destroy()

    def on_cancel(self):
        self.value = None
        self.destroy()


# ========================================================
#               PART 1: CONFIG & UTILS
# ========================================================

MODEL_URL = "https://github.com/thewh1teagle/kokoro-onnx/releases/download/model-files-v1.0/kokoro-v1.0.onnx"
VOICES_URL = "https://github.com/thewh1teagle/kokoro-onnx/releases/download/model-files-v1.0/voices-v1.0.bin"
MODEL_FILE = "kokoro-v1.0.onnx"
VOICES_FILE = "voices-v1.0.bin"


def clean_watermarks(text):
    text = re.sub(r'oceanofpdf\.com', '', text, flags=re.IGNORECASE)
    text = text.replace(' .', '.')
    return text


def download_file(url, filename, callback):
    if not os.path.exists(filename):
        callback(f"Downloading {filename}...")
        try:
            response = requests.get(url, stream=True)
            with open(filename, "wb") as f:
                for chunk in response.iter_content(chunk_size=8192):
                    if chunk: f.write(chunk)
            callback(f"Downloaded {filename}!")
        except Exception as e:
            callback(f"Download Error: {e}")


def extract_text_range(file_path, start_page=1, end_page=None):
    full_text = ""
    try:
        if file_path.endswith(".pdf"):
            doc = fitz.open(file_path)
            total_pages = len(doc)
            if end_page is None or end_page > total_pages:
                end_page = total_pages
            start_idx = max(0, start_page - 1)
            end_idx = min(total_pages, end_page)

            if start_idx >= total_pages:
                return f"Error: Start page {start_page} is beyond total pages ({total_pages})."

            text_list = []
            for i in range(start_idx, end_idx):
                text_list.append(doc[i].get_text())
            full_text = "\n".join(text_list)

        elif file_path.endswith(".docx"):
            doc = Document(file_path)
            full_text = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])

        return clean_watermarks(full_text)
    except Exception as e:
        return f"Error reading file: {e}"


def get_page_count(file_path):
    if file_path.endswith(".pdf"):
        try:
            doc = fitz.open(file_path)
            return len(doc)
        except:
            return 0
    return 0


# ========================================================
#               PART 2: THE APP (GUI)
# ========================================================

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")


class KokoroReader(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Book App (CPU Only)")
        self.geometry("1150x950")

        self.kokoro = None
        self.stop_event = threading.Event()
        self.pause_event = threading.Event()  # NEW: Pause control
        self.audio_queue = queue.Queue(maxsize=20)  # NEW: Expanded Buffer
        self.search_cursor = "1.0"
        self.current_file_path = None
        self.sel_start = "1.0"
        self.max_cpu_cores = os.cpu_count() or 4

        # --- HEADER ---
        self.frame_header = ctk.CTkFrame(self, fg_color="transparent")
        self.frame_header.pack(pady=15, padx=20, fill="x")

        self.btn_open = ctk.CTkButton(self.frame_header, text="📂 Open File", font=("Arial", 14, "bold"),
                                      width=140, height=40, command=self.open_file)
        self.btn_open.pack(side="left")

        self.voice_map = {
            "🇺🇸 Bella (Female)": "af_bella", "🇺🇸 Nicole (Female)": "af_nicole",
            "🇺🇸 Michael (Male)": "am_michael", "🇺🇸 Adam (Male)": "am_adam",
            "🇬🇧 Emma (Female)": "bf_emma", "🇬🇧 George (Male)": "bm_george"
        }
        self.voice_var = ctk.StringVar(value="🇺🇸 Michael (Male)")
        self.combo_voice = ctk.CTkComboBox(self.frame_header, values=list(self.voice_map.keys()),
                                           variable=self.voice_var, width=220, font=("Arial", 14))
        self.combo_voice.pack(side="right")
        ctk.CTkLabel(self.frame_header, text="Narrator:", font=("Arial", 14)).pack(side="right", padx=10)

        # --- SETTINGS FRAME ---
        self.frame_settings = ctk.CTkFrame(self, fg_color="#1e1e1e", corner_radius=10)
        self.frame_settings.pack(pady=5, padx=20, fill="x")

        # Speed Slider
        self.frame_speed = ctk.CTkFrame(self.frame_settings, fg_color="transparent")
        self.frame_speed.pack(side="left", fill="x", expand=True, padx=20, pady=10)
        ctk.CTkLabel(self.frame_speed, text="Reading Speed", font=("Arial", 12, "bold"), text_color="gray").pack(
            anchor="w")
        self.speed_slider = ctk.CTkSlider(self.frame_speed, from_=0.5, to=2.0, number_of_steps=15,
                                          command=self.update_speed_label)
        self.speed_slider.set(1.0)
        self.speed_slider.pack(fill="x", pady=5)
        self.lbl_speed_val = ctk.CTkLabel(self.frame_speed, text="1.0x", font=("Arial", 12, "bold"))
        self.lbl_speed_val.pack(anchor="e")

        ctk.CTkFrame(self.frame_settings, width=2, height=40, fg_color="#333333").pack(side="left", padx=5)

        # Threads Slider
        self.frame_threads = ctk.CTkFrame(self.frame_settings, fg_color="transparent")
        self.frame_threads.pack(side="right", fill="x", expand=True, padx=20, pady=10)
        ctk.CTkLabel(self.frame_threads, text="Process Threads (CPU)", font=("Arial", 12, "bold"),
                     text_color="gray").pack(
            anchor="w")
        self.thread_slider = ctk.CTkSlider(self.frame_threads, from_=1, to=self.max_cpu_cores,
                                           number_of_steps=self.max_cpu_cores - 1, command=self.update_thread_label)
        self.thread_slider.set(min(4, self.max_cpu_cores))
        self.thread_slider.pack(fill="x", pady=5)
        self.lbl_thread_val = ctk.CTkLabel(self.frame_threads, text=f"{int(self.thread_slider.get())} Threads",
                                           font=("Arial", 12, "bold"))
        self.lbl_thread_val.pack(anchor="e")

        # --- TEXT AREA ---
        self.text_frame = ctk.CTkFrame(self)
        self.text_frame.pack(pady=10, padx=20, fill="both", expand=True)

        self.text_box = tk.Text(self.text_frame, font=("Arial", 19), wrap="word",
                                bg="#2b2b2b", fg="#ffffff", selectbackground="#3b82f6",
                                insertbackground="white", bd=0, padx=20, pady=20)
        self.text_box.pack(fill="both", expand=True)
        self.text_box.tag_config("highlight", background="#facc15", foreground="black")

        # --- ACTIONS ---
        self.frame_actions = ctk.CTkFrame(self, height=80, fg_color="transparent")
        self.frame_actions.pack(pady=20, padx=20, fill="x")

        self.btn_read_sel = ctk.CTkButton(self.frame_actions, text="▶ Read Selection",
                                          fg_color="#10b981", hover_color="#059669", font=("Arial", 14, "bold"),
                                          command=self.start_reading_selection, width=150, height=50, corner_radius=25)
        self.btn_read_sel.pack(side="left", padx=(0, 10))

        self.btn_read_cur = ctk.CTkButton(self.frame_actions, text="▶ Read from Cursor",
                                          fg_color="#0ea5e9", hover_color="#0284c7", font=("Arial", 14, "bold"),
                                          command=self.start_reading_from_cursor, width=160, height=50,
                                          corner_radius=25)
        self.btn_read_cur.pack(side="left", padx=0)

        # NEW: Pause Button
        self.btn_pause = ctk.CTkButton(self.frame_actions, text="⏸ Pause",
                                       fg_color="#f59e0b", hover_color="#d97706", font=("Arial", 14, "bold"),
                                       command=self.toggle_pause, width=90, height=50, corner_radius=25,
                                       state="disabled")
        self.btn_pause.pack(side="left", padx=(15, 0))

        self.btn_stop = ctk.CTkButton(self.frame_actions, text="⏹ Stop",
                                      fg_color="#ef4444", hover_color="#dc2626", font=("Arial", 14, "bold"),
                                      command=self.stop_reading, width=90, height=50, corner_radius=25,
                                      state="disabled")
        self.btn_stop.pack(side="left", padx=10)

        self.btn_save = ctk.CTkButton(self.frame_actions, text="💾 Save to Audio",
                                      fg_color="#3b82f6", hover_color="#2563eb", font=("Arial", 14, "bold"),
                                      command=self.save_to_file, width=150, height=50, corner_radius=25)
        self.btn_save.pack(side="right", padx=0)

        # --- STATUS ---
        self.lbl_status = ctk.CTkLabel(self, text="Initializing Neural Engine...", text_color="#FCD34D",
                                       font=("Arial", 14, "bold"))
        self.lbl_status.pack(side="bottom", pady=(0, 10))

        self.progress_bar = ctk.CTkProgressBar(self, width=800, mode="determinate", progress_color="#10b981")
        self.progress_bar.set(0)

        providers = ort.get_available_providers()
        self.lbl_status.configure(text=f"Engine Mode: {providers[0]}", text_color="#FCD34D")

        if Kokoro is None:
            self.lbl_status.configure(text="Error: Kokoro Library Missing", text_color="red")
        else:
            threading.Thread(target=self.init_engine, daemon=True).start()

    def update_speed_label(self, value):
        self.lbl_speed_val.configure(text=f"{round(value, 2)}x")

    def update_thread_label(self, value):
        self.lbl_thread_val.configure(text=f"{int(value)} Threads")

    def init_engine(self):
        try:
            def update_status(msg):
                self.after(0, lambda: self.lbl_status.configure(text=msg))

            download_file(MODEL_URL, MODEL_FILE, update_status)
            download_file(VOICES_URL, VOICES_FILE, update_status)
            self.after(0, lambda: self.lbl_status.configure(text="Loading Model into Memory..."))

            self.kokoro = Kokoro(MODEL_FILE, VOICES_FILE)
            self.after(0, lambda: self.lbl_status.configure(text="Ready - Running on CPU", text_color="#10b981"))
        except Exception as e:
            self.after(0, lambda e=e: self.lbl_status.configure(text=f"Engine Load Failed: {e}", text_color="red"))

    def open_file(self):
        path = filedialog.askopenfilename(filetypes=[("Documents", "*.pdf *.docx")])
        if path:
            self.current_file_path = path
            threading.Thread(target=self.load_text_thread, args=(path,), daemon=True).start()

    def load_text_thread(self, path):
        self.after(0, lambda: self.lbl_status.configure(text="Extracting text from file..."))
        text = extract_text_range(path, start_page=1, end_page=None)
        if not text or not text.strip():
            text = "⚠️ Error: No text found."
        self.after(0, lambda: self.update_text_box(text))

    def update_text_box(self, text):
        self.text_box.delete("1.0", "end")
        self.text_box.insert("1.0", text)
        self.lbl_status.configure(text=f"Loaded: {os.path.basename(self.current_file_path)}")

    def split_into_sentences(self, text):
        text = text.replace("\n", " ")
        text = re.sub(r'\s+', ' ', text)
        sentences = re.split(r'(?<=[.!?])\s+', text)
        return [s.strip() for s in sentences if s.strip()]

    def start_reading_selection(self):
        try:
            self.sel_start = self.text_box.index("sel.first")
            self.sel_end = self.text_box.index("sel.last")
            text_block = self.text_box.get(self.sel_start, self.sel_end)
            self._start_reading_process(text_block, self.sel_start)
        except tk.TclError:
            ModernMessagebox(self, "No Selection", "Please highlight the text you want to read first!", "info")

    def start_reading_from_cursor(self):
        try:
            self.sel_start = self.text_box.index("insert")
            text_block = self.text_box.get(self.sel_start, "end")
            if not text_block.strip():
                ModernMessagebox(self, "End of Text", "There is no text after the cursor to read.", "info")
                return
            self._start_reading_process(text_block, self.sel_start)
        except Exception as e:
            print(f"Start Error: {e}")

    def _start_reading_process(self, text_block, start_pos):
        if not self.kokoro:
            ModernMessagebox(self, "Loading", "Engine is still loading, please wait...", "info")
            return

        self.stop_event.clear()
        self.pause_event.clear()

        self.btn_read_sel.configure(state="disabled")
        self.btn_read_cur.configure(state="disabled")
        self.btn_pause.configure(state="normal", text="⏸ Pause", fg_color="#f59e0b", hover_color="#d97706")
        self.btn_stop.configure(state="normal")

        while not self.audio_queue.empty():
            try:
                self.audio_queue.get_nowait()
            except:
                pass

        self.text_box.tag_remove("sel", "1.0", "end")
        self.text_box.tag_remove("highlight", "1.0", "end")

        self.search_cursor = start_pos
        current_speed = float(self.speed_slider.get())

        # Calculate total words for ETR
        total_words = len(text_block.split())

        threading.Thread(target=self.generator_thread, args=(text_block, current_speed), daemon=True).start()
        threading.Thread(target=self.playback_thread, args=(total_words, current_speed), daemon=True).start()

    def generator_thread(self, text_block, speed):
        sentences = self.split_into_sentences(text_block)
        voice_key = self.voice_map[self.voice_var.get()]

        for i, sentence in enumerate(sentences):
            if self.stop_event.is_set(): break
            try:
                clean_sentence = sentence.replace("\n", " ").strip()
                if not clean_sentence: continue

                audio, sample_rate = self.kokoro.create(clean_sentence, voice=voice_key, speed=speed, lang="en-us")
                padding = int(sample_rate * 0.25)
                silence = np.zeros(padding, dtype=np.float32)
                self.audio_queue.put((clean_sentence, np.concatenate((audio, silence)), sample_rate))
            except Exception as e:
                print(f"Gen Error on sentence: {e}")
        self.audio_queue.put(None)

    def toggle_pause(self):
        if self.pause_event.is_set():
            self.pause_event.clear()
            self.btn_pause.configure(text="⏸ Pause", fg_color="#f59e0b", hover_color="#d97706")
        else:
            self.pause_event.set()
            self.btn_pause.configure(text="▶ Resume", fg_color="#10b981", hover_color="#059669")

    def playback_thread(self, total_words, speed):
        self.after(0, lambda: self.lbl_status.configure(text="Buffering audio... please wait."))

        # Pre-buffering wait (smooths out CPU spikes)
        while self.audio_queue.qsize() < 4 and not self.stop_event.is_set():
            time.sleep(0.2)

        words_read = 0
        base_wpm = 150  # Standard reading speed
        current_wpm = base_wpm * speed

        while not self.stop_event.is_set():
            try:
                data = self.audio_queue.get(timeout=1)
                if data is None: break
                sentence, audio, sample_rate = data

                # Update UI and ETR
                words_read += len(sentence.split())
                words_remaining = max(0, total_words - words_read)
                etr_minutes = int(words_remaining / current_wpm)
                etr_seconds = int((words_remaining % current_wpm) / current_wpm * 60)

                status_text = f"Reading... | ETR: {etr_minutes}m {etr_seconds}s"

                self.after(0, lambda s=sentence: self.highlight_sentence(s))
                self.after(0, lambda st=status_text: self.lbl_status.configure(text=st))

                # --- NEW STREAMING PLAYBACK (Allows Instant Pause) ---
                try:
                    with sd.OutputStream(samplerate=sample_rate, channels=1) as stream:
                        chunk_size = 2048  # Small chunks of audio
                        for i in range(0, len(audio), chunk_size):
                            # Check pause state
                            while self.pause_event.is_set() and not self.stop_event.is_set():
                                time.sleep(0.1)

                            if self.stop_event.is_set(): break

                            # Write next chunk to speaker
                            chunk = np.float32(audio[i:i + chunk_size])
                            stream.write(chunk)
                except Exception as e:
                    print(f"Stream error: {e}")

            except queue.Empty:
                continue
            except Exception as e:
                break

        self.after(0, self.reset_ui)

    def highlight_sentence(self, sentence):
        count_var = tk.IntVar()
        pos_start = self.text_box.search(sentence, self.search_cursor, stopindex="end", count=count_var)
        if pos_start:
            pos_end = f"{pos_start} + {count_var.get()} chars"
            self.text_box.tag_remove("highlight", "1.0", "end")
            self.text_box.tag_add("highlight", pos_start, pos_end)
            self.text_box.see(pos_start)
            line_num = int(pos_start.split('.')[0])
            self.text_box.yview(line_num - 5 if line_num > 5 else 0)
            self.search_cursor = pos_end

    def stop_reading(self):
        self.stop_event.set()
        self.pause_event.clear()
        with self.audio_queue.mutex: self.audio_queue.queue.clear()
        self.reset_ui()

    def reset_ui(self):
        self.btn_read_sel.configure(state="normal")
        self.btn_read_cur.configure(state="normal")
        self.btn_pause.configure(state="disabled", text="⏸ Pause", fg_color="#f59e0b")
        self.btn_stop.configure(state="disabled")
        self.btn_save.configure(state="normal")
        self.text_box.tag_remove("highlight", "1.0", "end")
        self.lbl_status.configure(text="Ready")
        self.progress_bar.pack_forget()

    # --- SAVE TO FILE LOGIC (Unchanged) ---
    def save_to_file(self):
        if not self.current_file_path:
            ModernMessagebox(self, "Error", "Please open a file first!", "error")
            return

        total_pages = get_page_count(self.current_file_path)
        inp_start = ModernInputBox(self, "Start Page", f"Start Page (1 - {total_pages}):", "1")
        if not inp_start.value: return
        inp_end = ModernInputBox(self, "End Page", f"End Page (Total: {total_pages})", str(total_pages))

        try:
            start_num = int(inp_start.value)
            end_num = int(inp_end.value) if inp_end.value else total_pages
            if start_num < 1: raise ValueError
        except:
            ModernMessagebox(self, "Invalid Input", "Page numbers must be valid integers.", "error")
            return

        save_path = filedialog.asksaveasfilename(defaultextension=".wav",
                                                 filetypes=[("WAV Audio", "*.wav"), ("MP3 Audio", "*.mp3")])
        if not save_path: return

        self.btn_save.configure(state="disabled")
        self.btn_read_sel.configure(state="disabled")
        self.btn_read_cur.configure(state="disabled")
        self.progress_bar.pack(side="bottom", pady=5, padx=20)
        self.progress_bar.set(0)

        num_threads = int(self.thread_slider.get())
        self.lbl_status.configure(text=f"Turbo Mode: Converting with {num_threads} Threads...")

        threading.Thread(target=self.save_thread_parallel, args=(start_num, end_num, save_path, num_threads)).start()

    def generate_one_sentence(self, sentence, voice_key, speed):
        try:
            clean_sentence = sentence.replace("\n", " ").strip()
            if not clean_sentence: return None, 0
            audio, rate = self.kokoro.create(clean_sentence, voice=voice_key, speed=speed, lang="en-us")
            silence = np.zeros(int(rate * 0.25), dtype=np.float32)
            return np.concatenate((audio, silence)), rate
        except:
            return None, 0

    def save_thread_parallel(self, start_num, end_num, save_path, num_threads):
        try:
            full_text = extract_text_range(self.current_file_path, start_page=start_num, end_page=end_num)
            sentences = self.split_into_sentences(full_text)

            if not sentences:
                self.after(0, lambda: ModernMessagebox(self, "Error", "No text found to convert.", "error"))
                return

            voice_key = self.voice_map[self.voice_var.get()]
            speed = float(self.speed_slider.get())
            results_map = {}
            total = len(sentences)
            completed = 0
            start_time = time.time()

            with concurrent.futures.ThreadPoolExecutor(max_workers=num_threads) as executor:
                future_to_index = {executor.submit(self.generate_one_sentence, s, voice_key, speed): i for i, s in
                                   enumerate(sentences)}

                for future in concurrent.futures.as_completed(future_to_index):
                    idx = future_to_index[future]
                    audio, rate = future.result()
                    if audio is not None: results_map[idx] = (audio, rate)
                    completed += 1

                    elapsed = time.time() - start_time
                    if elapsed > 0 and completed > 0:
                        avg_per_item = elapsed / completed
                        remaining = total - completed
                        etr_seconds = int(remaining * avg_per_item)
                        mins, secs = divmod(etr_seconds, 60)
                        etr_str = f"{mins}m {secs}s"
                    else:
                        etr_str = "Calculating..."

                    progress = completed / total
                    self.after(0, lambda p=progress: self.progress_bar.set(p))
                    self.after(0, lambda c=completed, t=total, e=etr_str: self.lbl_status.configure(
                        text=f"Converted: {c}/{t} ({e} remaining)"))

            self.after(0, lambda: self.lbl_status.configure(text="Stitching & Saving file..."))
            ordered_audio = []
            final_rate = 24000
            for i in range(total):
                if i in results_map:
                    ordered_audio.append(results_map[i][0])
                    final_rate = results_map[i][1]

            if ordered_audio:
                final_audio = np.concatenate(ordered_audio)
                if save_path.endswith(".mp3"):
                    self.after(0, lambda: self.lbl_status.configure(text="Compressing to MP3 (This saves space!)..."))
                    audio_int16 = (final_audio * 32767).astype(np.int16)
                    song = AudioSegment(audio_int16.tobytes(), frame_rate=final_rate, sample_width=2, channels=1)
                    song.export(save_path, format="mp3", bitrate="128k")
                else:
                    sf.write(save_path, final_audio, final_rate)

                self.after(0, lambda: ModernMessagebox(self, "Success",
                                                       f"Audio Saved Successfully!\n\nLocation: {os.path.basename(save_path)}",
                                                       "success"))
            else:
                self.after(0, lambda: ModernMessagebox(self, "Error", "Generation Failed.", "error"))

        except Exception as e:
            self.after(0, lambda: ModernMessagebox(self, "Error",
                                                   f"Save Error: {e}\n\nIf using MP3, ensure FFmpeg is installed.",
                                                   "error"))
        finally:
            self.after(0, lambda: self.reset_ui())


if __name__ == "__main__":
    app = KokoroReader()
    app.mainloop()